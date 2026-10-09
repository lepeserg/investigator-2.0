"""Portable regression tests using only synthetic documents and mocked OCR."""
from pathlib import Path
import contextlib
import io
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile
import gc
import warnings
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'investigator-2-0/scripts'))
import word_text
import extract_docx
import ocr_smart
import make_docx

NS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def paragraph(text):return '<w:p><w:r><w:t>'+text+'</w:t></w:r></w:p>'
CHANGED='<w:p><w:del><w:r><w:delText>OLD_ONLY</w:delText></w:r></w:del><w:ins><w:r><w:t>NEW_ONLY</w:t></w:r></w:ins></w:p>'

class RefactorTests(unittest.TestCase):
    def setUp(self):
        (ROOT/'.test-tmp').mkdir(exist_ok=True)
        temp=tempfile.TemporaryDirectory(dir=ROOT/'.test-tmp')
        self.addCleanup(temp.cleanup);self.folder=Path(temp.name)

    def package(self,body,extra=None):
        f=self.folder/'synthetic.docx'
        with zipfile.ZipFile(f,'w') as z:
            z.writestr('word/document.xml',f'<w:document xmlns:w="{NS}"><w:body>{body}</w:body></w:document>')
            for part,xml in (extra or {}).items():z.writestr(part,xml)
        return f

    def test_plain_text_table_and_tail_keep_source_order(self):
        table='<w:tbl><w:tr><w:tc>'+paragraph('MIDDLE')+'</w:tc></w:tr></w:tbl>'
        f=self.package(paragraph('FIRST')+table+paragraph('LAST'))
        text,meta=word_text.read_docx(f)
        self.assertLess(text.index('FIRST'),text.index('MIDDLE'))
        self.assertLess(text.index('MIDDLE'),text.index('LAST'))
        self.assertEqual(meta['revisions'],{})

    def test_default_stops_before_returning_changed_content(self):
        f=self.package(CHANGED);before=f.read_bytes()
        with self.assertRaises(word_text.RevisionChoiceRequired):word_text.read_docx(f)
        self.assertEqual(before,f.read_bytes())

    def test_explicit_views_preserve_package_and_distinguish_versions(self):
        f=self.package(CHANGED);before=f.read_bytes()
        for view,kept,omitted in [('current','NEW_ONLY','OLD_ONLY'),('original','OLD_ONLY','NEW_ONLY')]:
            with self.subTest(view=view):
                text,meta=word_text.read_docx(f,revisions=view)
                self.assertIn(kept,text);self.assertNotIn(omitted,text);self.assertEqual(meta['view'],view)
        marked,_=word_text.read_docx(f,revisions='marked')
        self.assertIn('[ВСТАВКА:',marked);self.assertIn('[УДАЛЕНИЕ:',marked)
        self.assertEqual(before,f.read_bytes())

    def test_auxiliary_changes_block_even_body_only_reading(self):
        for part in ['header1','footer1','footnotes','endnotes']:
            with self.subTest(part=part):
                f=self.package(paragraph('BODY'),{f'word/{part}.xml':f'<w:hdr xmlns:w="{NS}">{CHANGED}</w:hdr>'})
                with self.assertRaises(word_text.RevisionChoiceRequired):word_text.read_docx(f)

    def test_formatting_revision_also_requires_choice(self):
        f=self.package('<w:p><w:pPr><w:pPrChange/></w:pPr><w:r><w:t>BODY</w:t></w:r></w:p>')
        with self.assertRaises(word_text.RevisionChoiceRequired):word_text.read_docx(f)

    def test_move_views_and_field_instruction(self):
        f=self.package('<w:p><w:moveFrom><w:r><w:t>FROM</w:t></w:r></w:moveFrom><w:moveTo><w:r><w:t>TO</w:t></w:r></w:moveTo><w:r><w:instrText>FIELD_CODE</w:instrText><w:t>VALUE</w:t></w:r></w:p>')
        current,_=word_text.read_docx(f,revisions='current');original,_=word_text.read_docx(f,revisions='original')
        self.assertIn('TO',current);self.assertNotIn('FROM',current);self.assertIn('FROM',original)
        self.assertNotIn('FIELD_CODE',current);self.assertIn('VALUE',current)

    def test_extract_cli_stops_without_printing_either_version(self):
        f=self.package(CHANGED)
        result=subprocess.run([sys.executable,str(ROOT/'investigator-2-0/scripts/extract_docx.py'),str(f)],capture_output=True,text=True,encoding='utf-8')
        self.assertNotEqual(result.returncode,0)
        self.assertNotIn('OLD_ONLY',result.stdout+result.stderr)
        self.assertNotIn('NEW_ONLY',result.stdout+result.stderr)

    def test_corpus_search_is_incomplete_until_explicit_choice(self):
        self.package(CHANGED)
        args=[sys.executable,str(ROOT/'investigator-2-0/scripts/corpus_search.py'),'NEW_ONLY','--root',str(self.folder)]
        result=subprocess.run(args,capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,2);self.assertIn('ПОИСК НЕПОЛНЫЙ',result.stdout+result.stderr)
        self.assertNotIn('В КОРПУСЕ НЕ НАЙДЕНО',result.stdout+result.stderr)
        chosen=subprocess.run(args+['--revisions','current'],capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(chosen.returncode,0,chosen.stderr);self.assertIn('NEW_ONLY',chosen.stdout)

    def test_ocr_nonzero_and_missing_output_are_errors(self):
        for code in [1,0]:
            with self.subTest(code=code),mock.patch.object(ocr_smart.subprocess,'run',return_value=subprocess.CompletedProcess([],code,b'',b'engine diagnostic')):
                with self.assertRaises(ocr_smart.OCRFailure):ocr_smart._recognize_page('tess','image',str(self.folder/'missing'),'PAGE 2')

    def test_ocr_successfully_empty_page_is_not_failure(self):
        base=self.folder/'blank';base.with_suffix('.txt').write_text('',encoding='utf-8')
        with mock.patch.object(ocr_smart.subprocess,'run',return_value=subprocess.CompletedProcess([],0,b'',b'')):
            self.assertEqual(ocr_smart._recognize_page('tess','image',str(base),'PAGE'), '')

    def test_ocr_cli_error_does_not_print_success_header_or_partial_text(self):
        image=self.folder/'scan.png';image.write_bytes(b'synthetic')
        stdout=io.StringIO();stderr=io.StringIO()
        with mock.patch.object(sys,'argv',['ocr_smart.py',str(image)]),mock.patch.object(ocr_smart,'find_marker',return_value=None),mock.patch.object(ocr_smart,'find_tesseract',return_value='tess'),mock.patch.object(ocr_smart,'run_tesseract',side_effect=ocr_smart.OCRFailure('PAGE 2 failed')),contextlib.redirect_stdout(stdout),contextlib.redirect_stderr(stderr):
            result=ocr_smart.main()
        self.assertEqual(result,1);self.assertNotIn('[OCR:',stdout.getvalue());self.assertIn('PAGE 2 failed',stderr.getvalue())

    def test_ocr_invalid_range_rejected_before_engine(self):
        with mock.patch.object(ocr_smart.subprocess,'run') as engine:
            with self.assertRaises(ValueError):ocr_smart.run_tesseract('tess','image.png',0,1)
            engine.assert_not_called()

    def test_pymupdf_supported_import_has_no_deprecation_warning(self):
        with warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter('always')
            with self.assertRaises(ValueError):ocr_smart.run_tesseract('tess','image.png',0,1)
            gc.collect()
        self.assertFalse([w for w in captured if 'deprecated' in str(w.message).lower()])

    def test_document_style_closes_local_constants_file(self):
        doc=Document()
        with warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter('always',ResourceWarning)
            make_docx._setup_document_style(doc)
            gc.collect()
        self.assertFalse([w for w in captured if issubclass(w.category,ResourceWarning)])

    def letterhead(self,valid=True):
        f=self.folder/'renamed-local-blank.docx';doc=Document();table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='УПРАВЛЕНИЕ ПО ЦЕНТРАЛЬНОМУ ВОЕННОМУ ОКРУГУ' if valid else 'SYNTHETIC ORGAN'
        for _ in range(4):table.cell(0,1).add_paragraph('')
        doc.save(f);return f

    def test_letterhead_recognized_after_rename_in_auto_and_explicit_modes(self):
        blank=self.letterhead();before=blank.read_bytes()
        for profile in ['auto','vsu-cvo-requests']:
            with self.subTest(profile=profile):
                out=self.folder/(profile+'.docx')
                make_docx.build_letterhead_doc(str(out),[{'type':'addressee','lines':['SYNTHETIC_RECIPIENT']},{'type':'para','text':'Учебный текст запроса.'}],blank_template_path=str(blank),letterhead_profile=profile)
                doc=Document(out)
                self.assertIn('SYNTHETIC_RECIPIENT',doc.tables[0].cell(0,1).text)
                self.assertNotIn('SYNTHETIC_RECIPIENT',doc.tables[0].cell(0,0).text)
        self.assertEqual(before,blank.read_bytes())

    def test_explicit_letterhead_requires_valid_structure(self):
        blank=self.letterhead(valid=False);out=self.folder/'rejected.docx'
        with self.assertRaises(ValueError):make_docx.build_letterhead_doc(str(out),[{'type':'addressee','lines':['SYNTHETIC_RECIPIENT']}],blank_template_path=str(blank),letterhead_profile='vsu-cvo-requests')
        self.assertFalse(out.exists())

    def test_letterhead_without_addressee_is_rejected(self):
        blank=self.letterhead();out=self.folder/'rejected.docx'
        with self.assertRaises(ValueError):make_docx.build_letterhead_doc(str(out),[{'type':'para','text':'Учебный текст.'}],blank_template_path=str(blank),letterhead_profile='vsu-cvo-requests')
        self.assertFalse(out.exists())

    def test_check_skill_flags_taskkill_by_image_name(self):
        import check_skill
        root=self.folder/'skill';(root/'references').mkdir(parents=True);(root/'scripts').mkdir()
        def flagged(code):
            (root/'scripts'/'synthetic.py').write_text('"""Синтетический скрипт."""\nimport os, subprocess\n'+code+'\n',encoding='utf-8')
            problems=check_skill.check(str(root))[0]
            return [p for p in problems if 'synthetic.py' in p and 'taskkill' in p]
        for code in ['os.system("taskkill /F /IM X.EXE")',
                     'os.system("TASKKILL /f /im excel.exe /T >nul 2>&1")',
                     'subprocess.run(["taskkill", "/F", "/IM", "X.EXE"])',
                     'subprocess.run([\n    "taskkill",\n    "/F",\n    "/im",\n    "WINWORD.EXE"])']:
            with self.subTest(code=code):self.assertTrue(flagged(code))
        for code in ['os.system("taskkill /F /PID 123")',
                     'subprocess.run(["taskkill", "/F", "/PID", str(123)])']:
            with self.subTest(code=code):self.assertFalse(flagged(code))

if __name__=='__main__':unittest.main()
