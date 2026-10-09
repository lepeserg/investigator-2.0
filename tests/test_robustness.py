"""Robustness: temp-file cleanup on failed saves, CLI usage."""
from pathlib import Path
import os
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT/'investigator-2-0/scripts'))
import docx_edit
import fix_statcard_font


def _replace_failing_for(target):
    """os.replace, падающий только на финальной замене target (doc_meta тоже зовёт os.replace)."""
    real = os.replace

    def fake(src, dst):
        if os.path.abspath(dst) == os.path.abspath(target):
            raise OSError('cross-device')
        return real(src, dst)
    return fake


def _leftovers(folder):
    return sorted(p.name for p in Path(folder).iterdir() if p.name.startswith('~'))


class SaveAtomicCleanupTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, 'doc.docx')
        doc = Document()
        doc.add_paragraph('исходный текст')
        doc.save(self.path)
        self.original = Path(self.path).read_bytes()

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_edit_cleans_tmp_when_save_fails(self):
        doc = Document(self.path)

        def broken_save(target):
            Path(target).write_bytes(b'partial')
            raise OSError('disk full')

        with patch.object(doc, 'save', side_effect=broken_save):
            with self.assertRaises(OSError):
                docx_edit._save_atomic(doc, self.path, backup=False)
        self.assertEqual(_leftovers(self.tmp.name), [])
        self.assertEqual(Path(self.path).read_bytes(), self.original)

    def test_docx_edit_cleans_tmp_when_replace_fails(self):
        doc = Document(self.path)
        with patch.object(docx_edit.os, 'replace', side_effect=_replace_failing_for(self.path)):
            with self.assertRaises(OSError):
                docx_edit._save_atomic(doc, self.path, backup=False)
        self.assertEqual(_leftovers(self.tmp.name), [])
        self.assertEqual(Path(self.path).read_bytes(), self.original)

    def test_docx_edit_success_leaves_no_tmp(self):
        doc = Document(self.path)
        doc.add_paragraph('новый абзац')
        docx_edit._save_atomic(doc, self.path, backup=False)
        self.assertEqual(_leftovers(self.tmp.name), [])
        self.assertIn('новый абзац', [p.text for p in Document(self.path).paragraphs])

    def test_statcard_font_cleans_tmp_when_save_fails(self):
        doc = Document(self.path)

        def broken_save(target):
            Path(target).write_bytes(b'partial')
            raise OSError('disk full')

        with patch.object(doc, 'save', side_effect=broken_save):
            with self.assertRaises(OSError):
                fix_statcard_font._save_atomic(doc, self.path)
        self.assertEqual(_leftovers(self.tmp.name), [])

    def test_statcard_font_cleans_tmp_when_replace_fails(self):
        doc = Document(self.path)
        with patch.object(fix_statcard_font.os, 'replace', side_effect=_replace_failing_for(self.path)):
            with self.assertRaises(OSError):
                fix_statcard_font._save_atomic(doc, self.path)
        self.assertEqual(_leftovers(self.tmp.name), [])
        self.assertEqual(Path(self.path).read_bytes(), self.original)


class StatcardsCliTests(unittest.TestCase):
    def test_missing_path_prints_usage_and_exits_2(self):
        script = ROOT/'investigator-2-0/scripts/make_statcards.py'
        for cmd in ('grid', 'verify', 'pdf'):
            result = subprocess.run([sys.executable, str(script), cmd], capture_output=True,
                                    text=True, encoding='utf-8', errors='replace', timeout=60)
            self.assertEqual(result.returncode, 2, cmd)
            self.assertIn('Не указан путь', result.stdout)
            self.assertNotIn('Traceback', result.stderr)


if __name__ == '__main__':
    unittest.main()
