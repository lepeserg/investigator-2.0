"""Synthetic document regression tests; no real case or legal conclusion is used."""
from pathlib import Path
import hashlib
import shutil
import sys
import tempfile
import unittest
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'investigator-2-0/scripts'))
import docx_edit
import docx_integrity
import make_docx
import check_tom
import consistency_check
import zipfile


class DocumentTests(unittest.TestCase):
    def setUp(self):
        (ROOT/'.test-tmp').mkdir(exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=ROOT/'.test-tmp')
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def test_working_copy_replaces_body_table_footer_and_preserves_source(self):
        source=self.root/'source.docx'; copy=self.root/'working.docx'
        doc=Document();doc.add_paragraph('SYNTHETIC_OLD subject')
        doc.add_table(rows=1,cols=2).cell(0,0).text='SYNTHETIC_OLD table'
        doc.sections[0].footer.paragraphs[0].text='SYNTHETIC_OLD footer'
        doc.save(source);before=source.read_bytes();shutil.copy2(source,copy)
        report=docx_edit.edit_file(str(copy),[{'op':'replace_all','old':'SYNTHETIC_OLD','new':'SYNTHETIC_NEW','expect':3}],backup=False)
        self.assertEqual(source.read_bytes(),before)
        after=Document(copy)
        self.assertIn('SYNTHETIC_NEW',after.paragraphs[0].text)
        self.assertIn('SYNTHETIC_NEW',after.tables[0].cell(0,0).text)
        self.assertIn('SYNTHETIC_NEW',after.sections[0].footer.paragraphs[0].text)
        self.assertEqual(len(after.tables),1)
        self.assertTrue(docx_integrity.check(str(copy)))

    def test_failed_preflight_leaves_file_unchanged(self):
        file=self.root/'draft.docx';doc=Document();doc.add_paragraph('SYNTHETIC TEXT');doc.save(file)
        before=file.read_bytes()
        with self.assertRaises(ValueError):
            docx_edit.edit_file(str(file),[{'op':'replace','old':'ABSENT TEXT','new':'NEW TEXT','expect':1}],backup=False)
        self.assertEqual(file.read_bytes(),before)

    def test_missing_letterhead_is_explicit_not_a_false_success(self):
        out=self.root/'request.docx'
        with self.assertRaises(ValueError):
            make_docx.build_letterhead_doc(str(out),[],blank_template_path=None)
        self.assertFalse(out.exists())



    def _docx_with_body(self, name, body_xml):
        """Synthetic DOCX whose document.xml body is replaced with raw WordprocessingML."""
        base=self.root/'base.docx';Document().save(base)
        out=self.root/name
        with zipfile.ZipFile(base) as src, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data=src.read(item.filename)
                if item.filename=='word/document.xml':
                    xml=data.decode('utf-8')
                    head=xml[:xml.index('<w:body>')+len('<w:body>')]
                    data=(head+body_xml+'</w:body></w:document>').encode('utf-8')
                dst.writestr(item,data)
        return out

    def test_consistency_text_skips_deleted_revisions_and_field_codes(self):
        body=('<w:p><w:r><w:t xml:space="preserve">Потерпевший </w:t></w:r>'
              '<w:del w:id="1" w:author="X" w:date="2026-01-01T00:00:00Z"><w:r>'
              '<w:delText>Старов А.А. 01.01.2020</w:delText></w:r></w:del>'
              '<w:ins w:id="2" w:author="X" w:date="2026-01-01T00:00:00Z"><w:r>'
              '<w:t>Новиков Б.Б.</w:t></w:r></w:ins></w:p>'
              '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
              '<w:r><w:instrText xml:space="preserve"> PAGE \\* MERGEFORMAT </w:instrText></w:r>'
              '<w:r><w:t>ООО &quot;Ромашка&quot; &amp; Ко</w:t></w:r></w:p>')
        text=consistency_check._docx_text(str(self._docx_with_body('tracked.docx',body)))
        self.assertIn('Новиков Б.Б.',text)
        self.assertNotIn('Старов',text)
        self.assertNotIn('01.01.2020',text)
        self.assertNotIn('MERGEFORMAT',text)
        self.assertIn('ООО "Ромашка" & Ко',text)

    def test_check_tom_fallback_ignores_tab_elements(self):
        xml=('<w:body><w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
             '<w:r><w:t>1.</w:t></w:r><w:r><w:tab/></w:r>'
             '<w:r><w:t xml:space="preserve">Протокол &amp; опись &quot;А&quot;</w:t></w:r></w:p></w:body>')
        self.assertEqual(check_tom._xml_paragraphs(xml),['1.Протокол & опись "А"'])

if __name__=='__main__':unittest.main()
