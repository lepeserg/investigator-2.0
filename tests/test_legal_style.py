"""Legal-style warnings on fictional fragments; no case material is used."""
from pathlib import Path
import hashlib
import subprocess
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'investigator-2-0' / 'scripts'))
import style_lint


class LegalStyleTests(unittest.TestCase):
    def setUp(self):
        (ROOT / '.test-tmp').mkdir(exist_ok=True)
        temp = tempfile.TemporaryDirectory(dir=ROOT / '.test-tmp')
        self.addCleanup(temp.cleanup)
        self.path = Path(temp.name) / 'synthetic.docx'
        doc = Document()
        doc.core_properties.author = 'Учебный пример'
        doc.core_properties.comments = 'Условные данные для проверки'
        doc.add_paragraph('Иванов И.И. и Петров П.П. встретились. Последний передал конверт.')
        doc.add_paragraph('Прошу представить сведения до 01.11.2026.')
        doc.add_paragraph('Прошу предоставить выписки и/или справки.')
        doc.save(self.path)

    def test_candidates_and_clean_fragments(self):
        cases = [
            ('Прошу предоставить выписки и/или справки.', {'стиль-проверить-перечень'}),
            ('Прошу предоставить выписки и справки.', set()),
            ('В записи сказано: «Предоставить выписки и/или справки до 01.11.2026».', set()),
            ('В записи сказано: "Предоставить справки до 01.11.2026".', set()),
            ('Прошу представить сведения до 01.11.2026.', {'стиль-проверить-границу'}),
            ('Прошу представить сведения до 01.11.2026 включительно.', set()),
            ('Прошу представить сведения до 01.11.2026 года включительно.', set()),
            ('Период до 01.11.2026, не включая эту дату.', set()),
            ('Проект до 31.02.2026.', set()),
            ('Иванов И.И. и Петров П.П. встретились. Последний передал конверт.', {'стиль-проверить-отсылку'}),
            ('Иванов И.И. и Петров П.П. встретились. Петров П.П. передал конверт.', set()),
            ('Иванов И.И. сообщил, что последний раз видел конверт у Петрова П.П.', set()),
            ('Иванов И.И. сообщил сведения. Последний передал конверт.', set()),
            ('Точную дату не помню, это было примерно в конце сентября.', set()),
            ('Об этом сообщил Иванов. Сам передачу не видел. Я испугался.', set()),
        ]
        for text, expected in cases:
            with self.subTest(text=text):
                self.assertEqual({item[0] for item in style_lint._legal_style_findings(text)}, expected)

    def test_opt_in_preserves_existing_results_and_file(self):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        baseline = style_lint.lint(str(self.path), genre='postanovlenie')
        reviewed = style_lint.lint(str(self.path), genre='postanovlenie', legal_style=True)
        self.assertFalse(any(item[1].startswith('стиль-проверить-') for item in baseline))
        self.assertEqual([item for item in reviewed if not item[1].startswith('стиль-проверить-')], baseline)
        self.assertEqual(sum(item[1].startswith('стиль-проверить-') for item in reviewed), 3)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def test_statcard_does_not_get_prose_warnings(self):
        findings = style_lint.lint(str(self.path), genre='statcard', legal_style=True)
        self.assertFalse(any(item[1].startswith('стиль-проверить-') for item in findings))

    def test_cli_reports_warnings_without_changing_docx(self):
        before = self.path.read_bytes()
        result = subprocess.run([sys.executable, str(ROOT / 'investigator-2-0' / 'scripts' / 'style_lint.py'),
                                 str(self.path), '--genre', 'postanovlenie', '--legal-style'],
                                capture_output=True, text=True, encoding='utf-8', errors='replace')
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertEqual(result.stdout.count('[стиль-проверить-'), 3)
        self.assertIn('кандидаты на сверку', result.stdout)
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
