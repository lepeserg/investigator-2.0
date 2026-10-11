"""Generate a separate readable Word report for each case-management operation."""
from datetime import datetime, timezone
from pathlib import Path
import json
import uuid

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from case_store import folder, safe_child


def document(title, intro, sections):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    for name, size in [('Normal', 12), ('Title', 19), ('Heading 1', 15), ('Heading 2', 13)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        fonts.set(qn('w:eastAsia'), 'Times New Roman')
        fonts.set(qn('w:cs'), 'Times New Roman')
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.05
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(intro)
    for heading, lines in sections:
        doc.add_heading(heading, 1)
        for index, line in enumerate(lines):
            paragraph = doc.add_paragraph(str(line))
            paragraph.paragraph_format.widow_control = True
            if len(lines) <= 5 and index < len(lines) - 1:
                paragraph.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = title
    return doc


def export(project, data, mode, sections):
    reports = safe_child(folder(project), 'reports')
    reports.mkdir(exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    stem = f'{stamp}-{mode}-{uuid.uuid4().hex[:8]}'
    path = safe_child(reports, stem + '.docx')
    title = 'Ведение дела ' + data['case']['label']
    intro = f"Рабочий отчёт по зарегистрированным сведениям. Редакция реестра {data['revision']}. Дата и область проверки указаны в разделах. Отчёт не подтверждает выполнение процессуальных действий."
    doc = document(title, intro, sections)
    doc.save(path)
    # Read back to check the generated package; page layout still requires rendering.
    Document(path)
    payload = {'revision': data['revision'], 'mode': mode, 'sections': sections,
               'word_file': str(path), 'visual_review': 'не выполнен программой'}
    safe_child(reports, stem + '.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return path
