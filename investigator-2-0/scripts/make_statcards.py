# -*- coding: utf-8 -*-
"""
Заполнение статистических карточек СК Ф-1 (на выявленное преступление) и
Ф-1.2 (на подозреваемого/обвиняемого).

Шаблоны (бланки Ф-1/Ф-1.2 с разметкой) — персональный слой (ДСП, вне поставки); путь передаётся аргументом, напр. references/templates/:
    Ф-1.docx     — карточка Ф-1
    Ф-1.2.docx   — карточка Ф-1.2  (внимание: иногда файл называют «Ф-12» —
                   это то же самое: Ф-1.2 «на подозреваемого (обвиняемого)»)
    ВНИМАНИЕ: в рабочей папке бланков владельца («Исходные бланки (с флешки)»)
    файл формы Ф-1.2 называется Ф-12.docx — точка в имени опущена. Демо ниже
    ищет «Ф-1.2.docx»; при запуске на этой папке имя надо передать своё.

Заполнение идёт по карте ячеек CELLMAP_F1 / CELLMAP_F12: ключ -> (таблица, строка, столбец).
Значение пишется в последний абзац ячейки (строка-значение под подсказкой формы),
шрифты и разметка бланка не трогаются. Если значение совпадает с уже имеющимся —
ячейка не изменяется (идемпотентность, эталон сохраняется байт-в-байт по тексту).

Использование:
    from make_statcards import fill_card, CELLMAP_F1, CELLMAP_F12
    fill_card(template_f1,  out_f1,  CELLMAP_F1,  DATA_F1)
    fill_card(template_f12, out_f12, CELLMAP_F12, DATA_F12)
"""
from docx import Document

# ---- Карта ячеек Ф-1 (на выявленное преступление) ----
CELLMAP_F1 = {
    "ud":                       (0, 4, 7),    # 1.1 номер УД (формат 1.ГГ.0200.2308.000XXX)
    "unit":                     (0, 5, 8),    # 1.2.а условное наим. в/части
    "crime_no":                 (0, 7, 16),   # 1.3 порядковый № преступления в деле (001)
    "krsp_no":                  (0, 8, 11),   # 1.4 номер КРСП
    "krsp_date":                (0, 8, 15),   # 1.4 дата КРСП (ДД.ММ.ГГ)
    "card_to_prosecutor_date":  (0, 9, 15),   # 1.5 дата направления карточки прокурору
    "commit_time_date":         (0, 15, 10),  # 2.3 время и дата совершения ("08: 30      ДД.ММ.ГГ")
    "vud_date":                 (0, 21, 14),  # 2.6.б дата возбуждения УД
    "accept_date":              (0, 22, 13),  # 2.6.в дата принятия к производству
    "opisanie":                 (0, 23, 4),   # 2.7 описание события (широкая ячейка справа; (0,23,0) — узкий столбец-номер!)
    "qualification":            (0, 26, 4),   # 2.8.а квалификация (ст. 337 ч. 5)
    "qual_features":            (0, 27, 4),   # 2.8.б квалифицирующие признаки (в период мобилизации)
    "region":                   (1, 13, 5),   # 2.14.б субъект РФ совершения
    "investigator":             (1, 18, 6),   # 2.19.а следователь (Фамилия И.О.)
    "invest_date":              (1, 18, 14),  # 2.19.а дата заполнения
    "leader":                   (1, 19, 5),   # 2.19.б руководитель СО
    "leader_date":              (1, 19, 14),  # 2.19.б дата
    "prosecutor":               (1, 20, 5),   # 2.19.в военный прокурор
}

# ---- Карта ячеек Ф-1.2 (на подозреваемого/обвиняемого) ----
CELLMAP_F12 = {
    "ud":                       (0, 3, 11),   # 1.1 номер УД
    "unit":                     (0, 4, 13),   # 1.2 условное наим. в/части
    "person_no":                (0, 5, 17),   # 1.3 порядковый № лица (001)
    "card_to_prosecutor_date":  (0, 6, 17),   # 1.4 дата направления карточки прокурору
    "surname":                  (0, 8, 1),    # 2.1.а фамилия
    "name":                     (0, 8, 6),    # 2.1.б имя
    "patronymic":               (0, 8, 14),   # 2.1.в отчество
    "position":                 (0, 9, 1),    # 2.2.а должность
    "rank":                     (0, 9, 10),   # 2.2.б звание
    "personal_no":              (0, 10, 15),  # 2.2.в личный номер
    "birth_date":               (0, 12, 17),  # 2.4.а дата рождения
    "birth_np":                 (0, 14, 4),   # 2.5.а место рождения — нас. пункт
    "birth_district":           (0, 15, 4),   # 2.5.б район
    "birth_region":             (0, 16, 4),   # 2.5.в область/край/государство
    "live_np":                  (0, 17, 5),   # 2.5.г-д место жительства — нас. пункт, район
    "live_region":              (0, 18, 5),   # 2.5.е область/край/государство
    "enlist_date":              (0, 23, 17),  # 2.8.в дата призыва / заключения контракта
    "nationality":              (0, 26, 4),   # 2.11 национальность
    "social_code":              (0, 27, 21),  # 2.12 код соц. положения (01 — в/сл по контракту)
    "position_code":            (0, 28, 21),  # 2.13 код должностного положения (43 — водитель)
    "crime_no2":                (0, 30, 7),   # 3.1.а № преступления (001)
    "qualification":            (0, 30, 8),   # 3.1.б квалификация (ст. 337 ч. 5)
    "vud_person_date":          (1, 0, 2),    # 3.2.а дата возбуждения УД в отношении лица
    "rozysk_date":              (1, 3, 11),   # 3.2.е дата объявления розыска
    "investigator":             (2, 0, 5),    # 3.4.а следователь
    "invest_date":              (2, 0, 7),    # 3.4.а дата
    "leader":                   (2, 1, 4),    # 3.4.б руководитель СО
    "leader_date":              (2, 1, 7),    # 3.4.б дата
    "prosecutor":               (2, 2, 4),    # 3.4.в военный прокурор
}


def _set_cell_value(cell, value):
    """Пишет значение в строку-значение ячейки (последний абзац), сохраняя
    шрифт и разметку бланка. Идемпотентно: если значение уже стоит — пропуск."""
    p = cell.paragraphs[-1]
    cur = "".join(r.text for r in p.runs)
    if cur.strip() == str(value).strip():
        return
    nonempty = [r for r in p.runs if r.text.strip()]
    if nonempty:
        nonempty[0].text = str(value)
        for r in nonempty[1:]:
            r.text = ""
    elif p.runs:
        p.runs[len(p.runs) // 2].text = str(value)
    else:
        p.add_run(str(value))


def fill_card(template_path, out_path, cellmap, data):
    doc = Document(template_path)
    for key, coords in cellmap.items():
        if key not in data:
            continue
        ti, r, c = coords
        _set_cell_value(doc.tables[ti].rows[r].cells[c], data[key])
    doc.save(out_path)
    return out_path


def _dump_grid(path):
    """grid: дамп всех непустых ячеек бланка как [t][r][c] -> текст (Контур 1 §09.00 п.2)."""
    import os
    doc = Document(path)
    print("Сетка ячеек: %s — %d таблиц(ы)" % (os.path.basename(path), len(doc.tables)))
    for ti, t in enumerate(doc.tables):
        print("\n== Таблица [%d]: %d строк x %d столбцов ==" % (ti, len(t.rows), len(t.columns)))
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                txt = cell.text.strip().replace("\n", " ")
                if txt:
                    print("  [%d][%d][%d] %s" % (ti, ri, ci, txt[:70]))


def _verify_map(path, which):
    """verify: показать, что стоит в каждой координате CELLMAP на ЭТОМ бланке (Контур 1)."""
    import os
    cellmap = CELLMAP_F12 if str(which).lower() in ("f12", "ф-1.2", "1.2", "12") else CELLMAP_F1
    name = "Ф-1.2" if cellmap is CELLMAP_F12 else "Ф-1"
    doc = Document(path)
    print("Сверка карты %s с бланком %s:" % (name, os.path.basename(path)))
    problems = 0
    for key, (ti, ri, ci) in cellmap.items():
        try:
            cur = doc.tables[ti].rows[ri].cells[ci].text.strip().replace("\n", " ")
            print("  %-24s -> [%d][%d][%d] = %s" % (key, ti, ri, ci, cur[:60]))
        except IndexError:
            print("  %-24s -> [%d][%d][%d] !! ВНЕ ДИАПАЗОНА" % (key, ti, ri, ci))
            problems += 1
    if problems:
        print("\n!! Карта НЕ совпадает с этим бланком: %d координат вне диапазона — снять свою карту через grid." % problems)
    else:
        print("\nВсе координаты попадают в таблицу. Значения сверить глазами по подписям граф формы.")


def _find_soffice():
    """PATH → типовые пути установки.

    ⚠ Одного `shutil.which` МАЛО: установщик LibreOffice на Windows себя в PATH не прописывает,
    поэтому рендер объявлялся недоступным при установленном LibreOffice (найдено 05.08.2026 —
    из-за этого в навык попала запись «LibreOffice НЕ установлен, рендер-сверка недоступна»).
    Тот же перебор уже был в check_env.find_libreoffice() — приводим к одному поведению."""
    import os, shutil
    for name in ("soffice", "soffice.exe", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/usr/bin/soffice", "/usr/bin/libreoffice",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(cand):
            return cand
    return None


def _render_pdf(path):
    """pdf: рендер .docx->PDF для визуальной сверки (Контур 2 §09.00 п.5). Нужен LibreOffice."""
    import os, subprocess
    soffice = _find_soffice()
    if not soffice:
        print("soffice/LibreOffice не найден: рендер .docx->PDF невозможен в этой среде "
              "(в Cowork soffice может отсутствовать). Варианты: открыть в Word локально; "
              "либо сверить вёрстку программно (grid/verify). pdftoppm .docx->PDF НЕ делает.")
        return 1
    outdir = os.path.dirname(os.path.abspath(path)) or "."
    out_pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    before = os.path.getmtime(out_pdf) if os.path.exists(out_pdf) else None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, path], check=False)
    # ⚠ soffice возвращает 0 даже когда ничего не сконвертировал (напр. занят другим процессом),
    # поэтому успех подтверждаем ПОЯВЛЕНИЕМ/ОБНОВЛЕНИЕМ файла, а не кодом возврата.
    if not os.path.exists(out_pdf) or (before is not None and os.path.getmtime(out_pdf) == before):
        print("soffice отработал, но PDF не появился/не обновился: %s\n"
              "Проверь, не открыт ли LibreOffice в другом окне, и права на папку." % out_pdf)
        return 1
    print("PDF готов: %s — открыть и сверить вёрстку/число страниц с донором." % out_pdf)
    return 0


# ============== CLI: grid / verify / pdf / demo ==============
if __name__ == "__main__":
    import os, sys, tempfile
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    _cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if _cmd in ("-h", "--help", "/?", "help"):
        print(__doc__)
        print("Подкоманды: grid | verify | pdf | demo — см. ниже.")
        print("  python make_statcards.py grid   <бланк.docx>")
        print("  python make_statcards.py verify <бланк.docx> [f1|f12]")
        print("  python make_statcards.py pdf    <карта.docx>")
        print("  python make_statcards.py demo   <out_dir> <tpl_dir>")
        sys.exit(0)
    if _cmd == "grid":
        _dump_grid(sys.argv[2]); sys.exit(0)
    if _cmd in ("verify", "check"):
        _verify_map(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "f1"); sys.exit(0)
    if _cmd == "pdf":
        sys.exit(_render_pdf(sys.argv[2]))
    if _cmd != "demo":
        print("Использование:\n"
              "  python make_statcards.py grid   <бланк.docx>            — дамп ячеек [t][r][c]->текст (Контур 1)\n"
              "  python make_statcards.py verify <бланк.docx> [f1|f12]   — сверить карту координат с бланком\n"
              "  python make_statcards.py pdf    <карта.docx>            — рендер в PDF (нужен LibreOffice, Контур 2)\n"
              "  python make_statcards.py demo   <out_dir> <tpl_dir>     — заполнить ДЕМО-данными (пример)\n"
              "Для заполнения реальной карты — fill_card(template, out, CELLMAP_F1/F12, DATA) из кода (см. docstring).")
        sys.exit(2)

    # Локальный пример исключён из публичной поставки.
    TPL = sys.argv[3] if len(sys.argv) > 3 else "{ПУТЬ_К_TEMPLATES}"
    OUT = sys.argv[2] if len(sys.argv) > 2 else tempfile.gettempdir()  # было "/tmp" — на Windows его нет

    DATA_F1 = {
        "ud": "УД № [номер]",
        "unit": "в/ч 00000",
        "crime_no": "[номер]",
        "krsp_no": "[номер КРСП]",
        "krsp_date": "ДД.ММ.ГГГГ",
        "card_to_prosecutor_date": "ДД.ММ.ГГГГ",
        "commit_time_date": "ДД.ММ.ГГГГ",
        "vud_date": "ДД.ММ.ГГГГ",
        "accept_date": "ДД.ММ.ГГГГ",
        "opisanie": "Условный учебный пример. Сведения о лице, событии и времени внесите по первоисточникам.",
        "qualification": "[квалификация]",
        "qual_features": "[признаки]",
        "region": "[регион]",
        "investigator": "[ФИО следователя]",
        "invest_date": "ДД.ММ.ГГГГ",
        "leader": "[ФИО руководителя]",
        "leader_date": "ДД.ММ.ГГГГ",
        "prosecutor": "[ФИО прокурора]",
    }
    DATA_F12 = {
        "ud": '1.26.0200.{КОД_ОРГАНА_УД}.000000',
        "unit": "00000",
        "person_no": "001",
        "card_to_prosecutor_date": "10.06.26",
        "surname": 'Фамилия', "name": "Иван", "patronymic": 'Фамилия',
        "position": "Водитель-электрик", "rank": "Рядовой",
        "personal_no": 'АВ-000000',
        "birth_date": "05.07.06",
        "birth_np": "Р.п. Пример", "birth_district": "Примерный р-н", "birth_region": "Примерная обл.",
        "live_np": "Р.п. Пример", "live_region": "Примерная обл.",
        "enlist_date": "19.06.25",
        "nationality": "Русский",
        "social_code": "01", "position_code": "43",
        "crime_no2": "001", "qualification": "ст. 337 ч. 5",
        "vud_person_date": "10.06.26", "rozysk_date": "10.06.26",
        "investigator": "{СЛЕДОВАТЕЛЬ_ФИО}", "invest_date": "10.06.26",
        "leader": "{РУКОВОДИТЕЛЬ_ФИО}", "leader_date": "10.06.26",
        "prosecutor": "{ПРОКУРОР_ФИО}",
    }
    p1 = fill_card(os.path.join(TPL, "Ф-1.docx"),   os.path.join(OUT, "Ф-1_test.docx"),   CELLMAP_F1,  DATA_F1)
    p2 = fill_card(os.path.join(TPL, "Ф-1.2.docx"), os.path.join(OUT, "Ф-1.2_test.docx"), CELLMAP_F12, DATA_F12)
    print("OK (демо):", p1, p2)
