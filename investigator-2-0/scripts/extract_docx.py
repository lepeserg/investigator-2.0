# -*- coding: utf-8 -*-
"""extract_docx.py — быстрый UTF-8 ридер содержимого .docx (абзацы + таблицы).

Зачем: читать содержимое .docx надёжно, без питон-однострочников (кириллица в
терминале ломается). Вывод всегда UTF-8.

Использование:
    python extract_docx.py "файл1.docx" "файл2.docx" ...
    python extract_docx.py "путь к папке"        # все .docx рекурсивно
    python extract_docx.py ... -o out.txt        # большой вывод в файл
"""
import os, sys, glob
os.environ.setdefault("PYTHONUTF8", "1")
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from docx import Document


def dump(path, out):
    """Возвращает True, если файл прочитан. False — если открыть не удалось
    (нужно для честного кода возврата при пакетном вызове, ревизия 22.08.2026)."""
    out.write("\n" + "=" * 80 + "\n")
    out.write("ФАЙЛ: " + os.path.basename(path) + "\n")
    out.write("=" * 80 + "\n")
    try:
        d = Document(path)
    except Exception as e:
        out.write("!! ошибка: %s\n" % e)
        return False
    for p in d.paragraphs:
        t = p.text.rstrip()
        if t:
            out.write(t + "\n")
    for ti, tbl in enumerate(d.tables):
        out.write("\n[ТАБЛИЦА %d]\n" % (ti + 1))
        for row in tbl.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            line = " | ".join(cells)
            if line.strip(" |"):
                out.write(line + "\n")
    return True


def collect(args):
    files, out_path = [], None
    i = 0
    while i < len(args):
        a = args[i]
        if a == "-o" and i + 1 < len(args):
            out_path = args[i + 1]
            i += 2
            continue
        if os.path.isdir(a):
            files += sorted(glob.glob(os.path.join(a, "**", "*.docx"), recursive=True))
        else:
            files.append(a)
        i += 1
    return files, out_path


def main(argv):
    if not argv:
        print("Укажи путь к .docx или папке. См. шапку extract_docx.py")
        return 2
    if argv[0] in ("-h", "--help", "/?", "help"):
        print(__doc__)
        return 0
    files, out_path = collect(argv)
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    bad = 0
    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            for p in files:
                if not dump(p, f):
                    bad += 1
        print("Записано %d файл(ов) в %s" % (len(files), out_path))
    else:
        for p in files:
            if not dump(p, sys.stdout):
                bad += 1
    if bad:
        # Раньше здесь возвращался 0 — «молчаливый успех» при непрочитанном файле.
        print("!! не удалось открыть файлов: %d из %d" % (bad, len(files)), file=sys.stderr)
        return 1
    if not files:
        print("Ни одного .docx не найдено.", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
