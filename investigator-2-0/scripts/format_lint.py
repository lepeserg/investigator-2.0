# -*- coding: utf-8 -*-
"""Проверить оформление Word-документа перед выдачей.

С --reference сравнивает форматную структуру с копией утверждённого донора.
Без образца проверяет общие параметры чистого листа. Визуальный просмотр
страниц остаётся необходимым в обоих режимах.
"""

import argparse
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_LINE_SPACING


_TEXT_TAGS = {
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}delText",
    "{http://schemas.openxmlformats.org/drawingml/2006/main}t",
}
_IGNORED_TAGS = {
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}proofErr",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lastRenderedPageBreak",
}


def _pt(value):
    return round(value.pt, 2) if value is not None else None


def _effective_font_size(paragraph, run, normal_size):
    if run.font.size is not None:
        return _pt(run.font.size)
    if paragraph.style.font.size is not None:
        return _pt(paragraph.style.font.size)
    return normal_size


def inspect(path, body_pt=13.0):
    doc = Document(path)
    issues = []
    normal = doc.styles["Normal"]
    normal_size = _pt(normal.font.size)
    if normal.font.name and normal.font.name.lower() != "times new roman":
        issues.append(f"основной стиль: шрифт {normal.font.name}, ожидается Times New Roman")
    if _pt(normal.font.size) != body_pt:
        issues.append(f"основной стиль: {_pt(normal.font.size)} пт, ожидается {body_pt:g} пт")
    spacing = normal.paragraph_format.line_spacing
    if spacing not in (1, 1.0) and normal.paragraph_format.line_spacing_rule != WD_LINE_SPACING.SINGLE:
        issues.append(f"основной стиль: межстрочный интервал {spacing!r}, ожидается одинарный")

    for i, sec in enumerate(doc.sections, 1):
        if abs(sec.page_width.cm - 21) > 0.05 or abs(sec.page_height.cm - 29.7) > 0.05:
            issues.append(f"раздел {i}: размер страницы {sec.page_width.cm:.2f} × {sec.page_height.cm:.2f} см, ожидается A4")

    paragraph_size_issues = {}
    for i, paragraph in enumerate(doc.paragraphs, 1):
        if not paragraph.text.strip():
            continue
        sizes = {_effective_font_size(paragraph, run, normal_size) for run in paragraph.runs if run.text.strip()}
        sizes.discard(None)
        allowed = {body_pt}
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER and paragraph.text.strip().isupper():
            allowed.add(14.0)  # название вида документа по §2.2
        if sizes and not sizes.issubset(allowed):
            paragraph_size_issues.setdefault(tuple(sorted(sizes)), []).append(i)
    for sizes, indices in paragraph_size_issues.items():
        issues.append(f"абзацы {indices}: размеры шрифта {list(sizes)} пт, ожидается {body_pt:g} пт")

    for ti, table in enumerate(doc.tables, 1):
        row_size_issues = {}
        for ri, row in enumerate(table.rows, 1):
            sizes = {
                _effective_font_size(paragraph, run, normal_size)
                for cell in row.cells
                for paragraph in cell.paragraphs
                for run in paragraph.runs
                if run.text.strip()
            }
            sizes.discard(None)
            if sizes and any(size != body_pt for size in sizes):
                row_size_issues.setdefault(tuple(sorted(sizes)), []).append(ri)
        for sizes, indices in row_size_issues.items():
            issues.append(f"таблица {ti}, строки {indices}: размеры шрифта {list(sizes)} пт, ожидается {body_pt:g} пт")
        if len(table.columns) >= 3 and len(table.rows) >= 2:
            widths = [col.width.cm if col.width else 0 for col in table.columns]
            lengths = [max(len(row.cells[ci].text) for row in table.rows) for ci in range(len(widths))]
            if widths and max(widths) - min(widths) < 0.1 and max(lengths) >= 3 * max(1, min(lengths)):
                issues.append(
                    f"таблица {ti}: равные ширины граф при разной длине данных "
                    f"({[round(x, 2) for x in widths]} см; максимум символов {lengths}); "
                    "проверить переносы на отрисованных страницах"
                )
    return issues


def _xml_signature(element):
    """Убрать значения текста, сохранив свойства, таблицы, табы и разрывы."""
    attrs = tuple(sorted((key, value) for key, value in element.attrib.items()
                         if not key.rsplit("}", 1)[-1].startswith("rsid")))
    children = tuple(_xml_signature(child) for child in element
                     if child.tag not in _IGNORED_TAGS)
    value = "" if element.tag in _TEXT_TAGS else (element.text or "").strip()
    return (element.tag, attrs, value, children)


def _format_parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        chosen = [name for name in names if name in {
            "word/document.xml", "word/styles.xml", "word/numbering.xml"
        } or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)]
        return {name: _xml_signature(ET.fromstring(archive.read(name))) for name in chosen}


def _first_difference(before, after, path=""):
    """Путь первого структурного отличия без вывода текста документа."""
    if before[:3] != after[:3] or len(before[3]) != len(after[3]):
        return path or "корень"
    for i, (old_child, new_child) in enumerate(zip(before[3], after[3]), 1):
        child_name = old_child[0].rsplit("}", 1)[-1]
        found = _first_difference(old_child, new_child, f"{path}/{child_name}[{i}]")
        if found:
            return found
    return ""


def compare_to_reference(path, reference):
    """Найти изменения форматной структуры после копирования донора.

    Добавленные абзацы и таблицы тоже сообщаются как отличия: их оформление
    проверяется вручную по отрисованным страницам.
    """
    doc = Document(path)
    base = Document(reference)
    issues = []
    if len(doc.sections) != len(base.sections):
        issues.append(f"разделов: образец {len(base.sections)}, результат {len(doc.sections)}")
    for i, (before, after) in enumerate(zip(base.sections, doc.sections), 1):
        for label, attr in (("ширина страницы", "page_width"), ("высота страницы", "page_height"),
                            ("левое поле", "left_margin"), ("правое поле", "right_margin"),
                            ("верхнее поле", "top_margin"), ("нижнее поле", "bottom_margin"),
                            ("верхний колонтитул", "header_distance"),
                            ("нижний колонтитул", "footer_distance")):
            old = getattr(before, attr)
            new = getattr(after, attr)
            if old != new:
                issues.append(f"раздел {i}, {label}: образец {old.cm:.2f} см, результат {new.cm:.2f} см")
    if len(base.tables) != len(doc.tables):
        issues.append(f"таблиц: образец {len(base.tables)}, результат {len(doc.tables)}")
    for i, (before, after) in enumerate(zip(base.tables, doc.tables), 1):
        old_shape = (len(before.rows), len(before.columns))
        new_shape = (len(after.rows), len(after.columns))
        if old_shape != new_shape:
            issues.append(f"таблица {i}: образец {old_shape}, результат {new_shape}")
        old_widths = [col.width for col in before.columns]
        new_widths = [col.width for col in after.columns]
        if old_widths != new_widths:
            issues.append(f"таблица {i}: изменилась ширина граф")
    old_parts = _format_parts(reference)
    new_parts = _format_parts(path)
    for name in sorted(old_parts.keys() | new_parts.keys()):
        if name not in old_parts:
            issues.append(f"добавлена форматная часть {name}")
        elif name not in new_parts:
            issues.append(f"пропала форматная часть {name}")
        elif old_parts[name] != new_parts[name]:
            where = _first_difference(old_parts[name], new_parts[name])
            issues.append(f"изменена форматная структура {name}, первое отличие {where} "
                          "(проверить табуляцию, отступы, шрифты и блоки)")
    return issues


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path, help="проверяемый .docx")
    parser.add_argument("--reference", type=Path, help="утверждённый донор .docx, от копии которого собран документ")
    parser.add_argument("--body-pt", type=float, default=13.0, help="размер основного шрифта по профилю (по умолчанию 13)")
    args = parser.parse_args()
    try:
        issues = compare_to_reference(args.file, args.reference) if args.reference else inspect(args.file, args.body_pt)
    except Exception as exc:
        print(f"Не удалось прочитать файл: {exc}", file=sys.stderr)
        return 2
    if issues:
        for issue in issues:
            print(f"ПРОВЕРИТЬ: {issue}")
        return 1
    print("format_lint: отклонений по проверяемым параметрам не найдено")
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    raise SystemExit(main())
