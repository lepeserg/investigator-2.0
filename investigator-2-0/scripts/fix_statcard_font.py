# -*- coding: utf-8 -*-
"""
fix_statcard_font.py — лечит главную болезнь статкарт: несовпадение шрифта ответов и формы.

Корень проблемы (диагностирован на бланках 384 ВСО): формы статкарт свёрстаны в Arial, а
ДЕФОЛТНЫЙ шрифт документа Word стоит Times New Roman. Поэтому любой ВПИСАННЫЙ ответ выходит
в Times New Roman (часто 12 пт) — не тем шрифтом и крупнее ячейки → «шрифт не нравится» +
«съехало» (строка растёт). Лечение: выставить дефолтный шрифт документа = шрифту формы и
(по флагу) перевести уже вписанные ответы в шрифт формы.

Запуск:
  python fix_statcard_font.py "<файл.docx>"                 # ОТЧЁТ (ничего не меняет)
  python fix_statcard_font.py "<файл.docx>" --apply         # выставить дефолт = шрифту формы
  python fix_statcard_font.py "<файл.docx>" --apply --convert   # + перевести ответы в шрифт формы
  python fix_statcard_font.py "<файл.docx>" --apply --target Arial --out "<новый.docx>"

Символьные шрифты (MT Extra, Wingdings, Symbol, Webdings) НЕ трогаются (иначе «галочки»/
формулы превратятся в мусор). Перед изменением исходника делается .bak.
"""
import sys, os, shutil, collections
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from docx import Document
from docx.oxml.ns import qn

SYMBOL_FONTS = {"MT Extra", "Wingdings", "Wingdings 2", "Wingdings 3", "Symbol", "Webdings"}


def _check_locked(path):
    """Единая проверка «открыт в Word» — берём из docx_edit, чтобы сообщение было
    одинаковым во всём навыке; если docx_edit недоступен, проверяем сами."""
    if not os.path.exists(path):
        return
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    try:
        import docx_edit
        docx_edit._check_locked(path)
        return
    except ImportError:
        pass
    d, name = os.path.split(path)
    d = d or "."
    for c in ("~$" + name, "~$" + name[2:] if len(name) >= 2 else "~$" + name):
        if os.path.exists(os.path.join(d, c)):
            sys.exit("Файл открыт в Word: %s. Закрой документ и повтори — правки НЕ внесены."
                     % os.path.basename(path))
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        sys.exit("Файл открыт в Word (или занят): %s. Закрой документ и повтори — "
                 "правки НЕ внесены." % os.path.basename(path))


def _save_atomic(doc, path):
    """Сохранить через tmp в ТОЙ ЖЕ папке -> проверка zip -> os.replace.
    Раньше python-docx писал прямо в оригинал: сбой посреди записи оставлял
    обрезанный .docx (ревизия 22.08.2026)."""
    import zipfile
    d = os.path.dirname(os.path.abspath(path)) or "."
    tmp = os.path.join(d, "~statcard_%d.docx" % os.getpid())
    k = 0
    while os.path.exists(tmp):
        k += 1
        tmp = os.path.join(d, "~statcard_%d_%d.docx" % (os.getpid(), k))
    doc.save(tmp)
    if not zipfile.is_zipfile(tmp):
        if os.path.exists(tmp):
            os.remove(tmp)
        sys.exit("python-docx сохранил битый файл — оригинал НЕ тронут.")
    try:
        os.replace(tmp, path)
    except PermissionError:
        if os.path.exists(tmp):
            os.remove(tmp)
        sys.exit("Не удалось записать «%s» — вероятно, открыт в Word. "
                 "Закрой документ и повтори — оригинал НЕ тронут." % os.path.basename(path))
    return path


def _all_runs(doc):
    def walk_paras(paras):
        for p in paras:
            for r in p.runs:
                yield r
    yield from walk_paras(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from walk_paras(cell.paragraphs)
                for nt in cell.tables:  # вложенные таблицы
                    for r2 in nt.rows:
                        for c2 in r2.cells:
                            yield from walk_paras(c2.paragraphs)


def font_profile(doc):
    c = collections.Counter()
    for r in _all_runs(doc):
        name = r.font.name
        if name and name not in SYMBOL_FONTS and r.text.strip():
            c[name] += 1
    return c


def _set_doc_default_font(doc, target):
    """docDefaults/rPrDefault/rPr/rFonts → target (ascii/hAnsi/cs)."""
    styles = doc.styles.element
    dd = styles.find(qn("w:docDefaults"))
    if dd is None:
        return False
    rprd = dd.find(qn("w:rPrDefault"))
    if rprd is None:
        rprd = dd.makeelement(qn("w:rPrDefault"), {}); dd.insert(0, rprd)
    rpr = rprd.find(qn("w:rPr"))
    if rpr is None:
        rpr = rprd.makeelement(qn("w:rPr"), {}); rprd.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {}); rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), target)
    return True


def _set_normal_font(doc, target):
    try:
        doc.styles["Normal"].font.name = target
        # cs тоже
        rpr = doc.styles["Normal"].element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            rf.set(qn("w:cs"), target)
        return True
    except Exception:
        return False


def _set_paragraph_mark_font(doc, target):
    """Шрифт МЕТКИ абзаца (pPr/rPr/rFonts) во всех абзацах, включая пустые ячейки таблиц.
    Именно из метки абзаца пустая графа наследует шрифт нового набора — без этого
    гарантия «type-safe» может не сработать (найдено аудитом 18.07.2026)."""
    def _paras(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from _paras(cell)
    n = 0
    for p in _paras(doc):
        pPr = p._p.get_or_add_pPr()
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = pPr.makeelement(qn("w:rPr"), {}); pPr.insert(0, rPr)
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = rPr.makeelement(qn("w:rFonts"), {}); rPr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), target)
        n += 1
    return n


def _convert_runs(doc, target):
    n = 0
    for r in _all_runs(doc):
        name = r.font.name
        if name in SYMBOL_FONTS:
            continue
        if name != target:
            r.font.name = target  # ставит ascii+hAnsi
            rpr = r._element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None:
                rf.set(qn("w:cs"), target)
            if r.text.strip():
                n += 1
    return n


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__); sys.exit(2)
    if args[0] in ("-h", "--help", "/?", "help"):
        print(__doc__); sys.exit(0)
    if args[0].startswith("-"):
        sys.exit("Первым аргументом нужен путь к .docx.\nСправка: python fix_statcard_font.py --help")
    path = args[0]
    if not os.path.isfile(path):
        sys.exit("Нет файла: %s" % path)
    apply = "--apply" in args
    convert = "--convert" in args
    target = None
    if "--target" in args:
        target = args[args.index("--target") + 1]
    out = None
    if "--out" in args:
        out = args[args.index("--out") + 1]

    doc = Document(path)
    prof = font_profile(doc)
    if not prof:
        print("Не нашёл текстовых рун с шрифтом — файл пустой или нестандартный.")
        sys.exit(2)
    form_font = target or prof.most_common(1)[0][0]

    print(f"Файл: {os.path.basename(path)}")
    print(f"Шрифты текста (без символьных): {dict(prof)}")
    print(f"Шрифт ФОРМЫ (доминирующий/целевой): {form_font}")
    # аномальные (не совпадающие с формой) руны — кандидаты в «ответы не тем шрифтом»
    anomalies = collections.Counter()
    big = []
    for r in _all_runs(doc):
        name = r.font.name
        if name and name not in SYMBOL_FONTS and name != form_font and r.text.strip():
            anomalies[name] += 1
            sz = r.font.size.pt if r.font.size else None
            if sz and sz >= 12:
                big.append((name, sz, r.text.strip()[:50]))
    dd_font = "—"
    styles = doc.styles.element.find(qn("w:docDefaults"))
    if styles is not None:
        rf = styles.find(qn("w:rPrDefault"))
        if rf is not None:
            rf = rf.find(qn("w:rPr"))
        if rf is not None:
            rf = rf.find(qn("w:rFonts"))
        if rf is not None:
            dd_font = rf.get(qn("w:ascii")) or "—"
    if dd_font == "—":
        note = "  (docDefaults не задан → текст наследует шрифт стиля Normal; если Normal = Times New Roman, ответы выйдут TNR)"
    elif dd_font != form_font:
        note = "  !! НЕ совпадает с формой — это и есть ловушка!"
    else:
        note = "  OK совпадает"
    print(f"Дефолтный шрифт документа (docDefaults): {dd_font}{note}")
    print(f"Ответов не в шрифте формы: {sum(anomalies.values())} рун {dict(anomalies)}")
    if big:
        print(f"Крупные (≥12 пт) вписанные значения — растят строку («съехало»): {len(big)}")
        for nm, sz, txt in big[:8]:
            print(f"   [{nm} {sz}] {txt}")

    if not apply:
        print("\n(ОТЧЁТ. Чтобы применить: добавьте --apply, для перевода ответов в шрифт формы — ещё --convert.)")
        sys.exit(0)

    # применяем
    target_path = out or path
    _check_locked(target_path)
    if not out:
        # .bak делается ОДИН раз — это первозданная карта до всех правок шрифта,
        # повторный прогон её не затирает (так и задумано).
        bak = path + ".bak"
        if not os.path.exists(bak):
            shutil.copy2(path, bak)
            print(f"\nБэкап: {os.path.basename(bak)}")
    d1 = _set_doc_default_font(doc, form_font)
    d2 = _set_normal_font(doc, form_font)
    marks = _set_paragraph_mark_font(doc, form_font)  # шрифт МЕТКИ абзаца в т.ч. пустых ячеек
    conv = _convert_runs(doc, form_font) if convert else 0
    _save_atomic(doc, target_path)
    print(f"\nГОТОВО → {target_path}")
    print(f"  docDefaults → {form_font}: {'ок' if d1 else 'нет docDefaults'}; Normal → {form_font}: {'ок' if d2 else 'нет'}; метки абзацев: {marks}")
    if convert:
        print(f"  переведено рун в {form_font}: {conv}")
    print(f"Новый вписанный текст должен идти шрифтом «{form_font}». ⚠ Проверить на СВОЁМ бланке: кликнуть в пустую графу и напечатать —")
    print("   если всё же выходит Times New Roman (ячейка форсирует свой шрифт), прогнать заполненную карту с --convert.")
    print("!! Обязательно сверить вёрстку: рендер в PDF (нужен LibreOffice) либо открыть в Word (в Cowork soffice может не быть).")


if __name__ == "__main__":
    main()
